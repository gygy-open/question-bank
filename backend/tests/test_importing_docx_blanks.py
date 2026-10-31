import zipfile
from pathlib import Path

from docx import Document

from app.services.importing.docx_blanks import preprocess_underline_blanks
from app.services.importing.ingest import DocxIngestor


def _build(path: Path, build) -> None:
    doc = Document()
    p = doc.add_paragraph()
    build(p)
    doc.save(path)


def _underlined(p, text: str) -> None:
    p.add_run(text).underline = True


def _convert(path: Path) -> str:
    import asyncio

    async def store(data: bytes, name: str):
        return None

    return asyncio.run(DocxIngestor().ingest(path, store=store)).markdown


def test_underlined_spaces_become_underscores(tmp_path):
    src = tmp_path / "a.docx"

    def build(p):
        p.add_run("工具是")
        _underlined(p, "         ")
        p.add_run("；")

    _build(src, build)
    assert "工具是\\_\\_\\_\\_\\_\\_\\_\\_\\_；" in _convert(src)


def test_adjacent_underlined_whitespace_runs_are_merged(tmp_path):
    src = tmp_path / "a.docx"

    def build(p):
        p.add_run("甲")
        _underlined(p, "  ")
        _underlined(p, "  ")
        p.add_run("乙")

    _build(src, build)
    assert "甲\\_\\_\\_\\_乙" in _convert(src)


def test_short_or_text_underline_is_left_alone(tmp_path):
    src = tmp_path / "a.docx"

    def build(p):
        p.add_run("甲")
        _underlined(p, " ")
        p.add_run("乙")
        _underlined(p, "重点")

    _build(src, build)
    out = tmp_path / "out.docx"
    assert preprocess_underline_blanks(src, out) == src
    assert not out.exists()


def test_underlined_tab_counts_as_blank(tmp_path):
    src = tmp_path / "a.docx"

    def build(p):
        p.add_run("甲")
        p.add_run("\t").underline = True
        p.add_run("乙")

    _build(src, build)
    out = tmp_path / "out.docx"
    assert preprocess_underline_blanks(src, out) == out
    with zipfile.ZipFile(out) as zf:
        assert "____".encode() in zf.read("word/document.xml")

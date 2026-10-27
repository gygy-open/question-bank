"""Excel 导入：多语言目录完整性、模板生成与跨语言表头解析。"""
import io
import re

import pytest
from openpyxl import Workbook, load_workbook

from app.core.locale import SUPPORTED_LOCALES, set_locale
from app.services import excel_import as xi
from app.services.excel_import import (
    EXCEL_EXAMPLES,
    EXCEL_TEXT,
    KP_SCHEMA,
    TAG_SCHEMA,
    USER_SCHEMA,
    MissingColumnsError,
)

SCHEMAS = [USER_SCHEMA, TAG_SCHEMA, KP_SCHEMA]
_PLACEHOLDER = re.compile(r"\{(\w+)\}")


@pytest.fixture(autouse=True)
def _reset_locale():
    yield
    set_locale("zh-CN")


def test_every_locale_has_the_same_keys():
    base = set(EXCEL_TEXT["zh-CN"])
    for loc in SUPPORTED_LOCALES:
        assert set(EXCEL_TEXT[loc]) == base, loc


def test_placeholders_match_across_locales():
    for key, zh in EXCEL_TEXT["zh-CN"].items():
        for loc in SUPPORTED_LOCALES:
            assert set(_PLACEHOLDER.findall(EXCEL_TEXT[loc][key])) == set(_PLACEHOLDER.findall(zh)), (loc, key)


@pytest.mark.parametrize("schema", SCHEMAS, ids=lambda s: s.name)
def test_every_column_has_a_header_and_examples_exist(schema):
    for loc in SUPPORTED_LOCALES:
        for col in schema.columns:
            assert f"{schema.name}.col.{col.key}" in EXCEL_TEXT[loc]
        assert EXCEL_EXAMPLES[schema.name][loc]
        for row in EXCEL_EXAMPLES[schema.name][loc]:
            assert len(row) <= len(schema.columns)


@pytest.mark.parametrize("schema", SCHEMAS, ids=lambda s: s.name)
def test_headers_do_not_collide_across_locales(schema):
    aliases = xi._aliases(schema)
    labels = {
        xi.header(schema, c.key, loc).strip().casefold()
        for loc in SUPPORTED_LOCALES
        for c in schema.columns
    }
    assert len(aliases) == len(labels)
    for loc in SUPPORTED_LOCALES:
        seen = [xi.header(schema, c.key, loc).casefold() for c in schema.columns]
        assert len(seen) == len(set(seen)), loc


@pytest.mark.parametrize("schema", SCHEMAS, ids=lambda s: s.name)
@pytest.mark.parametrize("loc", SUPPORTED_LOCALES)
def test_template_follows_locale_and_reads_back_in_any_locale(schema, loc):
    data = xi.generate_template(schema, loc)
    ws = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == [xi.header(schema, c.key, loc) for c in schema.columns]
    assert ws.title == xi.text(f"{schema.name}.sheet", loc)

    set_locale("en-US" if loc == "zh-CN" else "zh-CN")
    df = xi.read_rows(data, schema, schema.keys)
    assert list(df.columns) == schema.keys
    assert len(df) == len(EXCEL_EXAMPLES[schema.name][loc])


def test_missing_columns_reported_in_request_locale():
    wb = Workbook()
    wb.active.append(["Username"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(MissingColumnsError) as exc:
        xi.read_rows(buf.getvalue(), USER_SCHEMA)
    set_locale("en-US")
    msg = xi.missing_columns_message(exc.value)
    assert "Full Name" in msg and "Password" in msg and "Missing required columns" in msg
    set_locale("zh-CN")
    assert "缺少必需的列" in xi.missing_columns_message(exc.value)


def test_mixed_locale_headers_are_accepted():
    wb = Workbook()
    wb.active.append(["用户名", "Full Name", "Password"])
    wb.active.append(["u1", "A", "secret1"])
    buf = io.BytesIO()
    wb.save(buf)
    df = xi.read_rows(buf.getvalue(), USER_SCHEMA)
    assert list(df.columns) == ["username", "full_name", "password"]


def test_frontend_import_dialog_columns_match_template_headers():
    """导入对话框展示的列名必须和后端生成的模板表头一致（两边各维护一份文案）。"""
    import json
    from pathlib import Path

    from app.services import excel_import as xl

    locales = Path(__file__).resolve().parents[2] / "frontend" / "i18n" / "locales"
    expected = {
        "userImport": {
            "username": xl.header(xl.USER_SCHEMA, "username", "{loc}"),
            "fullName": xl.header(xl.USER_SCHEMA, "full_name", "{loc}"),
            "password": xl.header(xl.USER_SCHEMA, "password", "{loc}"),
            "admin": xl.header(xl.USER_SCHEMA, "admin", "{loc}"),
        },
        "tagImport": {
            "name": xl.header(xl.TAG_SCHEMA, "name", "{loc}"),
            "category": xl.header(xl.TAG_SCHEMA, "category", "{loc}"),
            "color": xl.header(xl.TAG_SCHEMA, "color", "{loc}"),
        },
        "kpImport": {
            "subject": xl.header(xl.KP_SCHEMA, "subject", "{loc}"),
            "level1": xl.header(xl.KP_SCHEMA, "level_1", "{loc}"),
            "level2": xl.header(xl.KP_SCHEMA, "level_2", "{loc}"),
            "level8": xl.header(xl.KP_SCHEMA, "level_8", "{loc}"),
        },
    }
    for loc in ("zh-CN", "en-US"):
        messages = json.loads((locales / f"{loc}.json").read_text(encoding="utf-8"))
        for ns, columns in expected.items():
            for name in columns:
                want = xl.header(
                    {"userImport": xl.USER_SCHEMA, "tagImport": xl.TAG_SCHEMA, "kpImport": xl.KP_SCHEMA}[ns],
                    {"fullName": "full_name", "level1": "level_1", "level2": "level_2", "level8": "level_8"}.get(name, name),
                    loc,
                )
                assert messages[ns]["columns"][name] == want, (loc, ns, name)

"""成绩录入第二条纵向切片的聚焦测试:名册 + 状态流转 + gradebook + 逐题录分。

覆盖:
- 学生/班级创建与权限
- 名单整体替换、去重、跨学科回滚(422)、班级 404
- 空名单不能开始录入;成功 start
- 逐题部分录入 total 为 NULL;完整录入合计
- 越界分数、重复题目 id 拒绝
- 乐观锁:同一 expected_revision 二次写 409
- batch_id 幂等:重试不重复事件、不二次加 revision
- 不同学生 revision 相互独立
- locked 后拒写;未录全不能 lock;录全可 lock
- 权限矩阵:viewer 无权看、editor 可看可录不可管理
- 跨 subject → 404
- 事件存在
"""
import io
import json
import uuid
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.security import create_access_token
from app.models.assessment import (
    Classroom,
    ClassroomStudent,
    ExamEvent,
    Student,
)
from app.models.question import Question, QuestionType
from app.models.subject import Subject
from app.models.subject_member import SubjectMember
from app.models.user import User

API = "/api/v1"


def _uid() -> str:
    return str(uuid.uuid4())


def _rich_doc(text: str) -> dict:
    return {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}]}


def _auth(user: User) -> dict:
    return {"Authorization": f"Bearer {create_access_token(subject=user.id)}"}


# --------------------------------------------------------------------------- #
# seed helpers
# --------------------------------------------------------------------------- #
async def _seed_user(db_session, *, username: str, is_superuser: bool = False) -> User:
    user = User(
        username=username, full_name=username, hashed_password="x",
        is_active=True, is_superuser=is_superuser,
    )
    db_session.add(user)
    await db_session.commit()
    await db_session.refresh(user)
    return user


async def _seed_subject(db_session, *, name="数学", slug="math") -> Subject:
    subject = Subject(name=name, slug=slug)
    db_session.add(subject)
    await db_session.commit()
    await db_session.refresh(subject)
    return subject


async def _grant(db_session, user, subject, role) -> None:
    db_session.add(SubjectMember(user_id=user.id, subject_id=subject.id, role=role))
    await db_session.commit()


async def _seed_full_question(db_session, *, subject_id: int, stem="1+1=?") -> Question:
    q = Question(
        content=json.dumps(_rich_doc(stem), ensure_ascii=False),
        options=[
            {"id": "opt_a", "label": "A", "content": _rich_doc("2")},
            {"id": "opt_b", "label": "B", "content": _rich_doc("3")},
        ],
        answer=json.dumps({"kind": "single_choice", "correct": "opt_a"}, ensure_ascii=False),
        analysis=json.dumps(_rich_doc("因为 1+1=2"), ensure_ascii=False),
        q_type=QuestionType.SINGLE_CHOICE,
        subject_id=subject_id,
        difficulty=3,
        source="seed",
        content_revision=1,
    )
    db_session.add(q)
    await db_session.commit()
    await db_session.refresh(q)
    return q


async def _seed_student(db_session, *, subject_id, student_no, name="学生") -> Student:
    student = Student(subject_id=subject_id, student_no=student_no, name=name)
    db_session.add(student)
    await db_session.commit()
    await db_session.refresh(student)
    return student


async def _seed_classroom(db_session, *, subject_id: int, name="一班", students=2) -> Classroom:
    classroom = Classroom(subject_id=subject_id, name=name)
    db_session.add(classroom)
    await db_session.commit()
    await db_session.refresh(classroom)
    for i in range(students):
        student = await _seed_student(
            db_session, subject_id=subject_id, student_no=f"{name}-{i:03d}", name=f"学生{i}"
        )
        db_session.add(ClassroomStudent(classroom_id=classroom.id, student_id=student.id))
    await db_session.commit()
    return classroom


# --------------------------------------------------------------------------- #
# 组稿 → 定稿(shared, scoring enabled)
# --------------------------------------------------------------------------- #
async def _finalized_version(client, sid, headers, *, q_ids, scores, scope="shared"):
    r = await client.post(
        f"{API}/subjects/{sid}/compositions?scope={scope}",
        json={"title": "稿件"}, headers=headers,
    )
    assert r.status_code == 201, r.text
    comp = r.json()
    r = await client.patch(
        f"{API}/subjects/{sid}/compositions/{comp['id']}?scope={scope}",
        json={"expected_revision": comp["revision"], "numbering_enabled": True, "scoring_enabled": True},
        headers=headers,
    )
    assert r.status_code == 200, r.text
    comp = r.json()
    rev = comp["revision"]
    nodes = []
    for qid, score in zip(q_ids, scores):
        props = {} if score is None else {"score": score}
        nodes.append({
            "id": _uid(), "node_kind": "block", "node_type": "question",
            "question_id": qid, "props": props,
        })
    r = await client.put(
        f"{API}/subjects/{sid}/compositions/{comp['id']}/nodes?scope={scope}",
        json={"expected_revision": rev, "nodes": nodes},
        headers=headers,
    )
    assert r.status_code == 200, r.text
    r = await client.post(
        f"{API}/subjects/{sid}/compositions/{comp['id']}/versions?scope={scope}",
        json={"expected_revision": rev + 1}, headers=headers,
    )
    assert r.status_code == 201, r.text
    return r.json()


async def _create_exam(client, sid, headers, *, version_id, classroom_id, name="期中考"):
    r = await client.post(
        f"{API}/subjects/{sid}/exam-sessions",
        json={"composition_version_id": version_id, "classroom_id": classroom_id, "name": name},
        headers=headers,
    )
    assert r.status_code == 201, r.text
    return r.json()


async def _make_recording_exam(client, db_session, ctx, *, students=2, scores=(40, 60)):
    """建 shared 定稿 + 班级 + 考试并进入 recording,返回 (exam, gradebook)。"""
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    qs = [await _seed_full_question(db_session, subject_id=sid, stem=f"Q{i}") for i in range(len(scores))]
    classroom = await _seed_classroom(db_session, subject_id=sid, students=students)
    version = await _finalized_version(client, sid, h, q_ids=[q.id for q in qs], scores=list(scores))
    exam = await _create_exam(client, sid, h, version_id=version["id"], classroom_id=classroom.id)
    r = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/start-recording", headers=h
    )
    assert r.status_code == 200, r.text
    gb = await client.get(f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/gradebook", headers=h)
    assert gb.status_code == 200, gb.text
    return exam, gb.json()


# --------------------------------------------------------------------------- #
# fixtures
# --------------------------------------------------------------------------- #
@pytest.fixture
async def ctx(db_session):
    admin = await _seed_user(db_session, username="admin", is_superuser=True)
    subject = await _seed_subject(db_session)
    subject2 = await _seed_subject(db_session, name="物理", slug="phys")

    editor = await _seed_user(db_session, username="editor")
    viewer = await _seed_user(db_session, username="viewer")
    manager = await _seed_user(db_session, username="manager")
    outsider = await _seed_user(db_session, username="outsider")
    for u, role in [(editor, "editor"), (viewer, "viewer"), (manager, "manager")]:
        await _grant(db_session, u, subject, role)
    await _grant(db_session, manager, subject2, "manager")

    return {
        "admin": admin, "subject": subject, "subject2": subject2,
        "editor": editor, "viewer": viewer, "manager": manager, "outsider": outsider,
    }


# --------------------------------------------------------------------------- #
# 学生 / 班级创建与权限
# --------------------------------------------------------------------------- #
async def test_create_student_and_list(client, ctx):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    r = await client.post(
        f"{API}/subjects/{sid}/students",
        json={"student_no": "S001", "name": "张三"}, headers=h,
    )
    assert r.status_code == 201, r.text
    assert r.json()["student_no"] == "S001"

    # 学号学科内唯一 → 冲突。
    r2 = await client.post(
        f"{API}/subjects/{sid}/students",
        json={"student_no": "S001", "name": "李四"}, headers=h,
    )
    assert r2.status_code == 409, r2.text

    lst = await client.get(f"{API}/subjects/{sid}/students?keyword=张", headers=h)
    assert lst.status_code == 200, lst.text
    body = lst.json()
    assert body["total"] == 1
    assert body["items"][0]["name"] == "张三"


async def test_create_student_binds_existing_user(client, ctx, db_session):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    bound = await _seed_user(db_session, username="pupil")

    r = await client.post(
        f"{API}/subjects/{sid}/students",
        json={"student_no": "S010", "name": "绑定生", "user_id": bound.id}, headers=h,
    )
    assert r.status_code == 201, r.text

    # 不存在的 user_id → 400。
    r2 = await client.post(
        f"{API}/subjects/{sid}/students",
        json={"student_no": "S011", "name": "错绑", "user_id": 999999}, headers=h,
    )
    assert r2.status_code == 400, r2.text

    # 同一 user 在同学科重复绑定 → 409。
    r3 = await client.post(
        f"{API}/subjects/{sid}/students",
        json={"student_no": "S012", "name": "重复绑", "user_id": bound.id}, headers=h,
    )
    assert r3.status_code == 409, r3.text


@pytest.mark.parametrize(
    "actor, expected",
    [("manager", 201), ("admin", 201), ("editor", 403), ("viewer", 403), ("outsider", 403)],
)
async def test_create_student_permission_matrix(client, ctx, actor, expected):
    sid = ctx["subject"].id
    r = await client.post(
        f"{API}/subjects/{sid}/students",
        json={"student_no": f"P-{actor}", "name": actor}, headers=_auth(ctx[actor]),
    )
    assert r.status_code == expected, r.text


@pytest.mark.parametrize(
    "actor, expected",
    [("manager", 201), ("admin", 201), ("editor", 403), ("viewer", 403), ("outsider", 403)],
)
async def test_create_classroom_permission_matrix(client, ctx, actor, expected):
    sid = ctx["subject"].id
    r = await client.post(
        f"{API}/subjects/{sid}/classrooms",
        json={"name": f"班-{actor}"}, headers=_auth(ctx[actor]),
    )
    assert r.status_code == expected, r.text


async def test_list_students_requires_view_assessment(client, ctx):
    sid = ctx["subject"].id
    r = await client.get(f"{API}/subjects/{sid}/students", headers=_auth(ctx["viewer"]))
    assert r.status_code == 403, r.text
    r = await client.get(f"{API}/subjects/{sid}/students", headers=_auth(ctx["editor"]))
    assert r.status_code == 200, r.text


# --------------------------------------------------------------------------- #
# 名单整体替换
# --------------------------------------------------------------------------- #
async def test_replace_classroom_members_dedup(client, ctx, db_session):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    classroom = await _seed_classroom(db_session, subject_id=sid, students=0)
    s1 = await _seed_student(db_session, subject_id=sid, student_no="A1")
    s2 = await _seed_student(db_session, subject_id=sid, student_no="A2")

    r = await client.put(
        f"{API}/subjects/{sid}/classrooms/{classroom.id}/students",
        json={"student_ids": [s1.id, s2.id, s1.id]}, headers=h,
    )
    assert r.status_code == 200, r.text
    assert len(r.json()) == 2

    # 整体替换为单人。
    r2 = await client.put(
        f"{API}/subjects/{sid}/classrooms/{classroom.id}/students",
        json={"student_ids": [s2.id]}, headers=h,
    )
    assert r2.status_code == 200, r2.text
    assert [s["id"] for s in r2.json()] == [s2.id]

    count = await db_session.scalar(
        select(func.count()).select_from(ClassroomStudent).where(
            ClassroomStudent.classroom_id == classroom.id
        )
    )
    assert count == 1


async def test_replace_members_cross_subject_rolls_back(client, ctx, db_session):
    sid = ctx["subject"].id
    sid2 = ctx["subject2"].id
    h = _auth(ctx["manager"])
    classroom = await _seed_classroom(db_session, subject_id=sid, students=0)
    good = await _seed_student(db_session, subject_id=sid, student_no="G1")
    foreign = await _seed_student(db_session, subject_id=sid2, student_no="F1")

    r = await client.put(
        f"{API}/subjects/{sid}/classrooms/{classroom.id}/students",
        json={"student_ids": [good.id, foreign.id]}, headers=h,
    )
    assert r.status_code == 422, r.text

    # 原子回滚:一个都没进班。
    count = await db_session.scalar(
        select(func.count()).select_from(ClassroomStudent).where(
            ClassroomStudent.classroom_id == classroom.id
        )
    )
    assert count == 0


async def test_replace_members_classroom_not_found(client, ctx):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    r = await client.put(
        f"{API}/subjects/{sid}/classrooms/999999/students",
        json={"student_ids": []}, headers=h,
    )
    assert r.status_code == 404, r.text


async def test_replace_members_requires_manage(client, ctx, db_session):
    sid = ctx["subject"].id
    classroom = await _seed_classroom(db_session, subject_id=sid, students=0)
    r = await client.put(
        f"{API}/subjects/{sid}/classrooms/{classroom.id}/students",
        json={"student_ids": []}, headers=_auth(ctx["editor"]),
    )
    assert r.status_code == 403, r.text


# --------------------------------------------------------------------------- #
# start-recording
# --------------------------------------------------------------------------- #
async def test_start_requires_participants(client, ctx, db_session):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    q = await _seed_full_question(db_session, subject_id=sid)
    empty = await _seed_classroom(db_session, subject_id=sid, students=0)
    version = await _finalized_version(client, sid, h, q_ids=[q.id], scores=[100])
    exam = await _create_exam(client, sid, h, version_id=version["id"], classroom_id=empty.id)

    r = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/start-recording", headers=h
    )
    assert r.status_code == 400, r.text


async def test_start_success_and_event(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=2)
    assert gb["status"] == "recording"

    started = await db_session.scalar(
        select(func.count()).select_from(ExamEvent).where(
            ExamEvent.exam_session_id == exam["id"],
            ExamEvent.event_type == "recording_started",
        )
    )
    assert started == 1


async def test_start_only_from_draft(client, ctx, db_session):
    exam, _ = await _make_recording_exam(client, db_session, ctx)
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    # 已 recording,再 start → 409。
    r = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/start-recording", headers=h
    )
    assert r.status_code == 409, r.text


@pytest.mark.parametrize("actor, expected", [("manager", 200), ("editor", 403), ("viewer", 403)])
async def test_start_permission_matrix(client, ctx, db_session, actor, expected):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    q = await _seed_full_question(db_session, subject_id=sid)
    classroom = await _seed_classroom(db_session, subject_id=sid, students=1)
    version = await _finalized_version(client, sid, h, q_ids=[q.id], scores=[100])
    exam = await _create_exam(client, sid, h, version_id=version["id"], classroom_id=classroom.id)

    r = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/start-recording",
        headers=_auth(ctx[actor]),
    )
    assert r.status_code == expected, r.text


# --------------------------------------------------------------------------- #
# 逐题录分:部分 / 完整 / total 计算
# --------------------------------------------------------------------------- #
def _row_for(gb, index=0):
    return gb["participants"][index]


def _qids(gb):
    return [q["id"] for q in gb["questions"]]


async def test_partial_scoring_keeps_total_null(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    row = _row_for(gb)
    qids = _qids(gb)

    r = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[0], "score": 30}]},
        headers=h,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total_score"] is None
    assert body["is_complete"] is False
    assert body["revision"] == row["revision"] + 1


async def test_full_scoring_sums_total(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    row = _row_for(gb)
    qids = _qids(gb)

    r = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={
            "expected_revision": row["revision"],
            "items": [
                {"exam_question_id": qids[0], "score": 35},
                {"exam_question_id": qids[1], "score": "55.50"},
            ],
        },
        headers=h,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert Decimal(str(body["total_score"])) == Decimal("90.50")
    assert body["is_complete"] is True

    # 清除一题 → total 回到 NULL。
    r2 = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": body["revision"], "items": [{"exam_question_id": qids[0], "score": None}]},
        headers=h,
    )
    assert r2.status_code == 200, r2.text
    assert r2.json()["total_score"] is None


# --------------------------------------------------------------------------- #
# 校验:越界 / 重复 id
# --------------------------------------------------------------------------- #
async def test_score_out_of_range_rejected(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    row = _row_for(gb)
    qids = _qids(gb)

    over = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[0], "score": 999}]},
        headers=h,
    )
    assert over.status_code in (400, 422), over.text

    neg = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[0], "score": -1}]},
        headers=h,
    )
    assert neg.status_code in (400, 422), neg.text


async def test_duplicate_question_id_rejected(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    row = _row_for(gb)
    qids = _qids(gb)

    r = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={
            "expected_revision": row["revision"],
            "items": [
                {"exam_question_id": qids[0], "score": 10},
                {"exam_question_id": qids[0], "score": 20},
            ],
        },
        headers=h,
    )
    assert r.status_code == 400, r.text


async def test_question_from_other_session_rejected(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    row = _row_for(gb)

    r = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": 999999, "score": 10}]},
        headers=h,
    )
    assert r.status_code == 422, r.text


# --------------------------------------------------------------------------- #
# 乐观锁 / 幂等 / 独立 revision
# --------------------------------------------------------------------------- #
async def test_stale_revision_conflict(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    row = _row_for(gb)
    qids = _qids(gb)

    ok = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[0], "score": 10}]},
        headers=h,
    )
    assert ok.status_code == 200, ok.text

    # 用旧 revision 再写 → 409。
    stale = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[1], "score": 20}]},
        headers=h,
    )
    assert stale.status_code == 409, stale.text


async def test_batch_id_idempotent(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    row = _row_for(gb)
    qids = _qids(gb)
    payload = {
        "expected_revision": row["revision"],
        "batch_id": "batch-xyz",
        "items": [{"exam_question_id": qids[0], "score": 12}],
    }
    url = f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}"

    first = await client.patch(url, json=payload, headers=h)
    assert first.status_code == 200, first.text
    assert first.json()["revision"] == row["revision"] + 1

    retry = await client.patch(url, json=payload, headers=h)
    assert retry.status_code == 200, retry.text
    # 幂等:revision 不再增加。
    assert retry.json()["revision"] == row["revision"] + 1

    events = await db_session.scalar(
        select(func.count()).select_from(ExamEvent).where(
            ExamEvent.exam_session_id == exam["id"],
            ExamEvent.batch_id == "batch-xyz",
        )
    )
    assert events == 1


async def test_batch_id_reuse_with_different_request_conflicts(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=2, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    first_row, second_row = gb["participants"]
    qids = _qids(gb)
    batch_id = "batch-reused"

    first = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{first_row['result_id']}",
        json={
            "expected_revision": first_row["revision"],
            "batch_id": batch_id,
            "items": [{"exam_question_id": qids[0], "score": 12}],
        },
        headers=h,
    )
    assert first.status_code == 200, first.text

    changed_payload = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{first_row['result_id']}",
        json={
            "expected_revision": first_row["revision"],
            "batch_id": batch_id,
            "items": [{"exam_question_id": qids[0], "score": 13}],
        },
        headers=h,
    )
    assert changed_payload.status_code == 409, changed_payload.text

    other_result = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{second_row['result_id']}",
        json={
            "expected_revision": second_row["revision"],
            "batch_id": batch_id,
            "items": [{"exam_question_id": qids[0], "score": 12}],
        },
        headers=h,
    )
    assert other_result.status_code == 409, other_result.text


async def test_independent_revisions_per_student(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=2, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    row0 = gb["participants"][0]
    row1 = gb["participants"][1]
    qids = _qids(gb)

    r = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row0['result_id']}",
        json={"expected_revision": row0["revision"], "items": [{"exam_question_id": qids[0], "score": 10}]},
        headers=h,
    )
    assert r.status_code == 200, r.text

    gb2 = (await client.get(f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/gradebook", headers=h)).json()
    by_id = {p["result_id"]: p for p in gb2["participants"]}
    assert by_id[row0["result_id"]]["revision"] == row0["revision"] + 1
    assert by_id[row1["result_id"]]["revision"] == row1["revision"]  # 未动


# --------------------------------------------------------------------------- #
# 锁定:未录全不能锁;录全可锁;锁后拒写
# --------------------------------------------------------------------------- #
async def _score_all(client, sid, exam_id, headers):
    """给所有参与者每题录满分。"""
    gb = (await client.get(f"{API}/subjects/{sid}/exam-sessions/{exam_id}/gradebook", headers=headers)).json()
    qmax = {q["id"]: q["max_score"] for q in gb["questions"]}
    for p in gb["participants"]:
        items = [{"exam_question_id": qid, "score": qmax[qid]} for qid in qmax]
        r = await client.patch(
            f"{API}/subjects/{sid}/exam-sessions/{exam_id}/results/{p['result_id']}",
            json={"expected_revision": p["revision"], "items": items}, headers=headers,
        )
        assert r.status_code == 200, r.text


async def test_lock_requires_all_complete(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=2, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    # 只给第一位录满,第二位空。
    p0 = gb["participants"][0]
    qmax = {q["id"]: q["max_score"] for q in gb["questions"]}
    await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{p0['result_id']}",
        json={"expected_revision": p0["revision"], "items": [{"exam_question_id": qid, "score": qmax[qid]} for qid in qmax]},
        headers=h,
    )

    r = await client.post(f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/lock", headers=h)
    assert r.status_code == 400, r.text


async def test_lock_success_and_write_after_lock_conflict(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=2, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    await _score_all(client, sid, exam["id"], h)

    r = await client.post(f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/lock", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "locked"

    locked_events = await db_session.scalar(
        select(func.count()).select_from(ExamEvent).where(
            ExamEvent.exam_session_id == exam["id"], ExamEvent.event_type == "locked"
        )
    )
    assert locked_events == 1

    # locked 后拒写 → 409。
    gb2 = (await client.get(f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/gradebook", headers=h)).json()
    row = gb2["participants"][0]
    qids = _qids(gb2)
    w = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[0], "score": 1}]},
        headers=h,
    )
    assert w.status_code == 409, w.text


async def test_lock_only_from_recording(client, ctx, db_session):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    q = await _seed_full_question(db_session, subject_id=sid)
    classroom = await _seed_classroom(db_session, subject_id=sid, students=1)
    version = await _finalized_version(client, sid, h, q_ids=[q.id], scores=[100])
    exam = await _create_exam(client, sid, h, version_id=version["id"], classroom_id=classroom.id)
    # draft 直接锁 → 409。
    r = await client.post(f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/lock", headers=h)
    assert r.status_code == 409, r.text


# --------------------------------------------------------------------------- #
# gradebook / save 权限矩阵
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    "actor, expected", [("manager", 200), ("editor", 200), ("viewer", 403), ("outsider", 403)]
)
async def test_gradebook_permission_matrix(client, ctx, db_session, actor, expected):
    exam, _ = await _make_recording_exam(client, db_session, ctx, students=1)
    sid = ctx["subject"].id
    r = await client.get(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/gradebook", headers=_auth(ctx[actor])
    )
    assert r.status_code == expected, r.text


async def test_editor_can_score_but_not_manage(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    row = _row_for(gb)
    qids = _qids(gb)

    # editor 可录分。
    r = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[0], "score": 10}]},
        headers=_auth(ctx["editor"]),
    )
    assert r.status_code == 200, r.text

    # editor 不能锁定(管理动作)。
    lock = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/lock", headers=_auth(ctx["editor"])
    )
    assert lock.status_code == 403, lock.text


async def test_viewer_cannot_save(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    row = _row_for(gb)
    qids = _qids(gb)
    r = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[0], "score": 10}]},
        headers=_auth(ctx["viewer"]),
    )
    assert r.status_code == 403, r.text


# --------------------------------------------------------------------------- #
# 跨 subject → 404
# --------------------------------------------------------------------------- #
async def test_cross_subject_gradebook_404(client, ctx, db_session):
    exam, _ = await _make_recording_exam(client, db_session, ctx, students=1)
    sid2 = ctx["subject2"].id
    h = _auth(ctx["manager"])  # manager 也管 subject2
    r = await client.get(
        f"{API}/subjects/{sid2}/exam-sessions/{exam['id']}/gradebook", headers=h
    )
    assert r.status_code == 404, r.text


async def test_cross_subject_save_404(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid2 = ctx["subject2"].id
    h = _auth(ctx["manager"])
    row = _row_for(gb)
    qids = _qids(gb)
    r = await client.patch(
        f"{API}/subjects/{sid2}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[0], "score": 10}]},
        headers=h,
    )
    assert r.status_code == 404, r.text


# --------------------------------------------------------------------------- #
# 事件存在
# --------------------------------------------------------------------------- #
async def test_scored_event_recorded(client, ctx, db_session):
    exam, gb = await _make_recording_exam(client, db_session, ctx, students=1, scores=(40, 60))
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    row = _row_for(gb)
    qids = _qids(gb)
    await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={"expected_revision": row["revision"], "items": [{"exam_question_id": qids[0], "score": 10}]},
        headers=h,
    )
    rows = (await db_session.execute(
        select(ExamEvent.event_type).where(ExamEvent.exam_session_id == exam["id"])
    )).scalars().all()
    assert "created" in rows
    assert "recording_started" in rows
    assert "scores_saved" in rows


# --------------------------------------------------------------------------- #
# Excel 导出 / 预检 / 原子应用
# --------------------------------------------------------------------------- #
async def _export_gradebook(client, sid, exam_id, headers) -> bytes:
    response = await client.get(
        f"{API}/subjects/{sid}/exam-sessions/{exam_id}/gradebook.xlsx",
        headers=headers,
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return response.content


def _set_import_scores(content: bytes, rows: list[list[float]]) -> bytes:
    workbook = load_workbook(io.BytesIO(content))
    worksheet = workbook["成绩录入"]
    for row_offset, scores in enumerate(rows, start=2):
        for column_offset, score in enumerate(scores, start=3):
            worksheet.cell(row=row_offset, column=column_offset, value=score)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


async def test_gradebook_excel_preview_apply_and_idempotency(client, ctx, db_session):
    exam, gradebook = await _make_recording_exam(
        client, db_session, ctx, students=2, scores=(40, 60)
    )
    sid = ctx["subject"].id
    headers = _auth(ctx["manager"])
    content = await _export_gradebook(client, sid, exam["id"], headers)

    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["成绩录入", "_meta"]
    assert workbook["_meta"].sheet_state == "hidden"
    assert workbook["成绩录入"].cell(row=2, column=1).value == gradebook["participants"][0]["student_no"]

    edited = _set_import_scores(content, [[30, 50], [35, 55]])
    files = {
        "file": (
            "scores.xlsx",
            edited,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    preview = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/score-imports/preview",
        files=files,
        headers=headers,
    )
    assert preview.status_code == 200, preview.text
    assert preview.json() == {
        "valid": True,
        "changed_rows": 2,
        "changed_scores": 4,
        "unchanged_rows": 0,
        "errors": [],
    }

    # 预检只读，不增加 result revision。
    before = (
        await client.get(
            f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/gradebook",
            headers=headers,
        )
    ).json()
    assert [row["revision"] for row in before["participants"]] == [1, 1]

    apply_files = {"file": files["file"]}
    applied = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/score-imports/apply",
        data={"batch_id": "excel-batch-1"},
        files=apply_files,
        headers=headers,
    )
    assert applied.status_code == 200, applied.text
    assert applied.json() == {
        "batch_id": "excel-batch-1",
        "updated_rows": 2,
        "updated_scores": 4,
    }

    # 完全相同的文件和 batch_id 重放，不重复增加 revision / 事件。
    retry = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/score-imports/apply",
        data={"batch_id": "excel-batch-1"},
        files={"file": files["file"]},
        headers=headers,
    )
    assert retry.status_code == 200, retry.text
    after = (
        await client.get(
            f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/gradebook",
            headers=headers,
        )
    ).json()
    assert [row["revision"] for row in after["participants"]] == [2, 2]
    assert [Decimal(row["total_score"]) for row in after["participants"]] == [
        Decimal("80"),
        Decimal("90"),
    ]


async def test_gradebook_excel_invalid_score_is_atomic(client, ctx, db_session):
    exam, _ = await _make_recording_exam(client, db_session, ctx, students=2, scores=(40, 60))
    sid = ctx["subject"].id
    headers = _auth(ctx["manager"])
    content = await _export_gradebook(client, sid, exam["id"], headers)
    edited = _set_import_scores(content, [[30, 50], [999, 55]])
    files = {"file": ("scores.xlsx", edited, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    preview = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/score-imports/preview",
        files=files,
        headers=headers,
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["valid"] is False
    assert preview.json()["errors"][0]["row"] == 3

    applied = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/score-imports/apply",
        data={"batch_id": "excel-invalid"},
        files={"file": files["file"]},
        headers=headers,
    )
    assert applied.status_code == 422, applied.text

    unchanged = (
        await client.get(
            f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/gradebook",
            headers=headers,
        )
    ).json()
    assert [row["revision"] for row in unchanged["participants"]] == [1, 1]
    assert all(row["total_score"] is None for row in unchanged["participants"])


async def test_gradebook_excel_permissions(client, ctx, db_session):
    exam, _ = await _make_recording_exam(client, db_session, ctx, students=1)
    sid = ctx["subject"].id

    exported = await client.get(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/gradebook.xlsx",
        headers=_auth(ctx["editor"]),
    )
    assert exported.status_code == 200, exported.text

    denied_export = await client.get(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/gradebook.xlsx",
        headers=_auth(ctx["viewer"]),
    )
    assert denied_export.status_code == 403, denied_export.text

    files = {"file": ("scores.xlsx", exported.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/score-imports/preview",
        files=files,
        headers=_auth(ctx["editor"]),
    )
    assert preview.status_code == 200, preview.text

    denied_preview = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/score-imports/preview",
        files={"file": files["file"]},
        headers=_auth(ctx["viewer"]),
    )
    assert denied_preview.status_code == 403, denied_preview.text


async def test_gradebook_excel_stale_template_rejected(client, ctx, db_session):
    exam, gradebook = await _make_recording_exam(
        client, db_session, ctx, students=1, scores=(40, 60)
    )
    sid = ctx["subject"].id
    headers = _auth(ctx["manager"])
    exported = await _export_gradebook(client, sid, exam["id"], headers)
    row = gradebook["participants"][0]
    first_question = gradebook["questions"][0]

    saved = await client.patch(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/results/{row['result_id']}",
        json={
            "expected_revision": row["revision"],
            "items": [{"exam_question_id": first_question["id"], "score": 10}],
        },
        headers=headers,
    )
    assert saved.status_code == 200, saved.text

    stale = _set_import_scores(exported, [[20, 30]])
    preview = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{exam['id']}/score-imports/preview",
        files={"file": ("stale.xlsx", stale, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["valid"] is False
    assert "重新导出" in preview.json()["errors"][0]["message"]


# --------------------------------------------------------------------------- #
# 班级成员读取:GET /classrooms/{id}/students
# --------------------------------------------------------------------------- #
async def test_list_classroom_students_returns_members(client, ctx, db_session):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    classroom = await _seed_classroom(db_session, subject_id=sid, name="三班", students=3)
    r = await client.get(
        f"{API}/subjects/{sid}/classrooms/{classroom.id}/students", headers=h
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert len(body) == 3
    assert {s["student_no"] for s in body} == {"三班-000", "三班-001", "三班-002"}
    assert all(s["subject_id"] == sid for s in body)


async def test_list_classroom_students_empty(client, ctx, db_session):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    classroom = await _seed_classroom(db_session, subject_id=sid, students=0)
    r = await client.get(
        f"{API}/subjects/{sid}/classrooms/{classroom.id}/students", headers=h
    )
    assert r.status_code == 200, r.text
    assert r.json() == []


async def test_list_classroom_students_requires_view_assessment(client, ctx, db_session):
    sid = ctx["subject"].id
    classroom = await _seed_classroom(db_session, subject_id=sid, students=1)
    r = await client.get(
        f"{API}/subjects/{sid}/classrooms/{classroom.id}/students",
        headers=_auth(ctx["viewer"]),
    )
    assert r.status_code == 403, r.text
    r = await client.get(
        f"{API}/subjects/{sid}/classrooms/{classroom.id}/students",
        headers=_auth(ctx["editor"]),
    )
    assert r.status_code == 200, r.text


async def test_list_classroom_students_cross_subject_404(client, ctx, db_session):
    sid = ctx["subject"].id
    sid2 = ctx["subject2"].id
    h = _auth(ctx["manager"])  # manager 也管 subject2
    classroom = await _seed_classroom(db_session, subject_id=sid, students=1)
    r = await client.get(
        f"{API}/subjects/{sid2}/classrooms/{classroom.id}/students", headers=h
    )
    assert r.status_code == 404, r.text


async def test_list_classroom_students_missing_404(client, ctx):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    r = await client.get(f"{API}/subjects/{sid}/classrooms/999999/students", headers=h)
    assert r.status_code == 404, r.text


# --------------------------------------------------------------------------- #
# 考试列表读取:GET /exam-sessions
# --------------------------------------------------------------------------- #
async def test_list_exam_sessions_scoped_and_ordered(client, ctx, db_session):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    q = await _seed_full_question(db_session, subject_id=sid)
    classroom = await _seed_classroom(db_session, subject_id=sid, students=1)
    version = await _finalized_version(client, sid, h, q_ids=[q.id], scores=[100])
    e1 = await _create_exam(client, sid, h, version_id=version["id"], classroom_id=classroom.id, name="第一场")
    e2 = await _create_exam(client, sid, h, version_id=version["id"], classroom_id=classroom.id, name="第二场")

    # subject2 下另建一场,不应出现在 subject 的列表里。
    sid2 = ctx["subject2"].id
    q2 = await _seed_full_question(db_session, subject_id=sid2)
    classroom2 = await _seed_classroom(db_session, subject_id=sid2, name="外班", students=1)
    version2 = await _finalized_version(client, sid2, h, q_ids=[q2.id], scores=[100])
    await _create_exam(client, sid2, h, version_id=version2["id"], classroom_id=classroom2.id, name="外场")

    r = await client.get(f"{API}/subjects/{sid}/exam-sessions", headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    ids = [row["id"] for row in body]
    assert set(ids) == {e1["id"], e2["id"]}
    # created_at desc:后建的在前(id 兜底倒序)。
    assert ids[0] == e2["id"]
    assert all("questions" not in row for row in body)


async def test_list_exam_sessions_status_and_classroom_filter(client, ctx, db_session):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    q = await _seed_full_question(db_session, subject_id=sid)
    room_a = await _seed_classroom(db_session, subject_id=sid, name="A班", students=1)
    room_b = await _seed_classroom(db_session, subject_id=sid, name="B班", students=1)
    version = await _finalized_version(client, sid, h, q_ids=[q.id], scores=[100])
    draft = await _create_exam(client, sid, h, version_id=version["id"], classroom_id=room_a.id, name="草稿场")
    recording = await _create_exam(client, sid, h, version_id=version["id"], classroom_id=room_b.id, name="录入场")
    r = await client.post(
        f"{API}/subjects/{sid}/exam-sessions/{recording['id']}/start-recording", headers=h
    )
    assert r.status_code == 200, r.text

    # status 过滤。
    r = await client.get(f"{API}/subjects/{sid}/exam-sessions?status=recording", headers=h)
    assert r.status_code == 200, r.text
    assert [row["id"] for row in r.json()] == [recording["id"]]

    r = await client.get(f"{API}/subjects/{sid}/exam-sessions?status=draft", headers=h)
    assert [row["id"] for row in r.json()] == [draft["id"]]

    # classroom_id 过滤。
    r = await client.get(f"{API}/subjects/{sid}/exam-sessions?classroom_id={room_a.id}", headers=h)
    assert [row["id"] for row in r.json()] == [draft["id"]]


async def test_list_exam_sessions_requires_view_assessment(client, ctx, db_session):
    sid = ctx["subject"].id
    r = await client.get(f"{API}/subjects/{sid}/exam-sessions", headers=_auth(ctx["viewer"]))
    assert r.status_code == 403, r.text
    r = await client.get(f"{API}/subjects/{sid}/exam-sessions", headers=_auth(ctx["editor"]))
    assert r.status_code == 200, r.text


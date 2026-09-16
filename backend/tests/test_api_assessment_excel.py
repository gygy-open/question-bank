"""通用评测成绩册 Excel v2 的纵向测试。

覆盖:
- 导出:`_meta` format/version/session、条目列、行映射;当前最新评分入册、未评分留空;
- 预览:合法(无变更)/ 有变更统计 / 行错误聚合 / stale attempt revision;
- 应用:追加评分历史 + attempt 缓存刷新;错误文件原子拒绝;
  batch+sha256 幂等重放;同 batch 不同文件 409;权限矩阵;跨学科 404。

评测/名册经真实 API + db_session 注入,成绩册读写用 openpyxl 直接操作工作簿。
"""
import io
import json
import uuid
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.security import create_access_token
from app.models.assessment import ResponseGrade
from app.models.question import Question, QuestionType
from app.models.subject import Subject
from app.models.subject_member import SubjectMember
from app.models.user import User
from app.services.assessment_excel import FORMAT, FORMAT_VERSION, META_SHEET_NAME, SHEET_NAME

API = "/api/v1"


def _uid() -> str:
    return str(uuid.uuid4())


def _rich_doc(text: str) -> dict:
    return {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}]}


def _auth(user: User) -> dict:
    return {"Authorization": f"Bearer {create_access_token(subject=user.id)}"}


# --------------------------------------------------------------------------- #
# seed helpers
# --------------------------------------------------------------------------- #
async def _seed_user(db_session, *, username: str, is_superuser: bool = False) -> User:
    user = User(
        username=username, full_name=username, hashed_password="x",
        is_active=True, is_superuser=is_superuser,
    )
    db_session.add(user)
    await db_session.commit()
    await db_session.refresh(user)
    return user


async def _seed_subject(db_session, *, name="数学", slug="math") -> Subject:
    subject = Subject(name=name, slug=slug)
    db_session.add(subject)
    await db_session.commit()
    await db_session.refresh(subject)
    return subject


async def _grant(db_session, user, subject, role) -> None:
    db_session.add(SubjectMember(user_id=user.id, subject_id=subject.id, role=role))
    await db_session.commit()


async def _seed_question(db_session, *, subject_id: int, stem="1+1=?") -> Question:
    q = Question(
        content=json.dumps(_rich_doc(stem), ensure_ascii=False),
        options=[
            {"id": "opt_a", "label": "A", "content": _rich_doc("2")},
            {"id": "opt_b", "label": "B", "content": _rich_doc("3")},
        ],
        answer=json.dumps({"kind": "single_choice", "correct": "opt_a"}, ensure_ascii=False),
        analysis=json.dumps(_rich_doc("因为 1+1=2"), ensure_ascii=False),
        q_type=QuestionType.SINGLE_CHOICE,
        subject_id=subject_id,
        difficulty=3,
        source="seed",
        content_revision=1,
    )
    db_session.add(q)
    await db_session.commit()
    await db_session.refresh(q)
    return q


async def _seed_classroom(db_session, *, subject_id: int, name="一班", students=2):
    from app.models.assessment import Classroom, ClassroomStudent, Student

    classroom = Classroom(subject_id=subject_id, name=name)
    db_session.add(classroom)
    await db_session.commit()
    await db_session.refresh(classroom)
    for i in range(students):
        student = Student(subject_id=subject_id, student_no=f"{name}-{i:03d}", name=f"学生{i}")
        db_session.add(student)
        await db_session.commit()
        await db_session.refresh(student)
        db_session.add(ClassroomStudent(classroom_id=classroom.id, student_id=student.id))
    await db_session.commit()
    return classroom


async def _finalized_version(client, sid, headers, *, q_ids, scores, scope="shared"):
    r = await client.post(
        f"{API}/subjects/{sid}/compositions?scope={scope}",
        json={"title": "稿件"}, headers=headers,
    )
    assert r.status_code == 201, r.text
    comp = r.json()
    r = await client.patch(
        f"{API}/subjects/{sid}/compositions/{comp['id']}?scope={scope}",
        json={"expected_revision": comp["revision"], "numbering_enabled": True, "scoring_enabled": True},
        headers=headers,
    )
    assert r.status_code == 200, r.text
    comp = r.json()
    rev = comp["revision"]
    nodes = [
        {
            "id": _uid(), "node_kind": "block", "node_type": "question",
            "question_id": qid, "props": ({} if score is None else {"score": score}),
        }
        for qid, score in zip(q_ids, scores)
    ]
    r = await client.put(
        f"{API}/subjects/{sid}/compositions/{comp['id']}/nodes?scope={scope}",
        json={"expected_revision": rev, "nodes": nodes}, headers=headers,
    )
    assert r.status_code == 200, r.text
    r = await client.post(
        f"{API}/subjects/{sid}/compositions/{comp['id']}/versions?scope={scope}",
        json={"expected_revision": rev + 1}, headers=headers,
    )
    assert r.status_code == 201, r.text
    return r.json()


# --------------------------------------------------------------------------- #
# fixtures
# --------------------------------------------------------------------------- #
@pytest.fixture
async def ctx(db_session):
    admin = await _seed_user(db_session, username="admin", is_superuser=True)
    subject = await _seed_subject(db_session)
    subject2 = await _seed_subject(db_session, name="物理", slug="phys")
    editor = await _seed_user(db_session, username="editor")
    viewer = await _seed_user(db_session, username="viewer")
    manager = await _seed_user(db_session, username="manager")
    for u, role in [(editor, "editor"), (viewer, "viewer"), (manager, "manager")]:
        await _grant(db_session, u, subject, role)
    await _grant(db_session, manager, subject2, "manager")
    return {
        "admin": admin, "subject": subject, "subject2": subject2,
        "editor": editor, "viewer": viewer, "manager": manager,
    }


async def _make_session(client, ctx, db_session, *, students=2, scores=(40, 60), title="期中评测"):
    sid = ctx["subject"].id
    h = _auth(ctx["manager"])
    qs = [await _seed_question(db_session, subject_id=sid, stem=f"Q{i}") for i in range(len(scores))]
    classroom = await _seed_classroom(db_session, subject_id=sid, students=students)
    version = await _finalized_version(client, sid, h, q_ids=[q.id for q in qs], scores=list(scores))
    r = await client.post(
        f"{API}/subjects/{sid}/assessments",
        json={"title": title, "composition_version_id": version["id"], "classroom_id": classroom.id},
        headers=h,
    )
    assert r.status_code == 201, r.text
    return sid, h, qs, version, r.json()


async def _grade_attempt(client, sid, headers, attempt, item_scores, *, expected_revision=None):
    body = {
        "expected_revision": attempt["revision"] if expected_revision is None else expected_revision,
        "items": [{"item_id": iid, "score": s} for iid, s in item_scores],
    }
    return await client.patch(
        f"{API}/subjects/{sid}/assessment-attempts/{attempt['attempt_id']}/grades",
        json=body, headers=headers,
    )


async def _start(client, sid, session_id, headers):
    r = await client.post(
        f"{API}/subjects/{sid}/assessment-sessions/{session_id}/grading/start", headers=headers
    )
    assert r.status_code == 200, r.text


async def _download(client, sid, session_id, headers):
    r = await client.get(
        f"{API}/subjects/{sid}/assessment-sessions/{session_id}/gradebook.xlsx", headers=headers
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return r.content


# --------------------------------------------------------------------------- #
# workbook helpers
# --------------------------------------------------------------------------- #
def _meta_maps(content: bytes):
    """返回 (item_id→列号, attempt_id→行号, head dict)。"""
    wb = load_workbook(io.BytesIO(content))
    meta = wb[META_SHEET_NAME]
    rows = list(meta.iter_rows(values_only=True))
    head, item_col, row_by_attempt = {}, {}, {}
    section = None
    for r in rows:
        if not r:
            continue
        if r[0] in ("format", "version", "session_id"):
            head[r[0]] = r[1]
            continue
        if r[0] == "items":
            section = "items_h"; continue
        if r[0] == "rows":
            section = "rows_h"; continue
        if section == "items_h":
            section = "items"; continue
        if section == "rows_h":
            section = "rows"; continue
        if section == "items":
            item_col[int(r[1])] = int(r[0])
        elif section == "rows" and r[2] is not None:
            row_by_attempt[int(r[2])] = int(r[0])
    return item_col, row_by_attempt, head


def _set_cells(content: bytes, edits):
    """edits: list of (attempt_id, item_id, value)。返回改写后的工作簿字节。"""
    wb = load_workbook(io.BytesIO(content))
    item_col, row_by_attempt, _ = _meta_maps(content)
    sheet = wb[SHEET_NAME]
    for attempt_id, item_id, value in edits:
        sheet.cell(row_by_attempt[attempt_id], item_col[item_id]).value = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _upload_files(content: bytes):
    return {"file": ("gradebook.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


async def _preview(client, sid, session_id, headers, content):
    return await client.post(
        f"{API}/subjects/{sid}/assessment-sessions/{session_id}/grade-imports/preview",
        files=_upload_files(content), headers=headers,
    )


async def _apply(client, sid, session_id, headers, content, *, batch_id=None):
    data = {"batch_id": batch_id} if batch_id is not None else None
    return await client.post(
        f"{API}/subjects/{sid}/assessment-sessions/{session_id}/grade-imports/apply",
        files=_upload_files(content), data=data, headers=headers,
    )


# --------------------------------------------------------------------------- #
# 导出:元数据 + 当前评分
# --------------------------------------------------------------------------- #
async def test_export_metadata_and_current_grade(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    items = session["items"]
    attempt = session["participants"][0]["attempts"][0]
    # 只评第一题:第一题入册,第二题留空。
    r = await _grade_attempt(client, sid, h, attempt, [(items[0]["id"], 40)])
    assert r.status_code == 200, r.text

    content = await _download(client, sid, session["id"], h)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames[:1] == [SHEET_NAME]
    assert wb[META_SHEET_NAME].sheet_state == "hidden"

    item_col, row_by_attempt, head = _meta_maps(content)
    assert head["format"] == FORMAT
    assert int(head["version"]) == FORMAT_VERSION
    assert int(head["session_id"]) == session["id"]
    assert set(item_col) == {items[0]["id"], items[1]["id"]}
    assert attempt["attempt_id"] in row_by_attempt

    sheet = wb[SHEET_NAME]
    row = row_by_attempt[attempt["attempt_id"]]
    assert Decimal(str(sheet.cell(row, item_col[items[0]["id"]]).value)) == Decimal("40")
    assert sheet.cell(row, item_col[items[1]["id"]]).value is None  # 未评分留空


# --------------------------------------------------------------------------- #
# 预览:合法 / 变更 / 错误聚合 / stale
# --------------------------------------------------------------------------- #
async def test_preview_valid_no_change(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    content = await _download(client, sid, session["id"], h)
    r = await _preview(client, sid, session["id"], h, content)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["valid"] is True
    assert body["changed_rows"] == 0
    assert body["changed_scores"] == 0
    assert body["errors"] == []


async def test_preview_reports_changes(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, students=2, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    items = session["items"]
    a0 = session["participants"][0]["attempts"][0]
    content = await _download(client, sid, session["id"], h)
    edited = _set_cells(content, [(a0["attempt_id"], items[0]["id"], 40), (a0["attempt_id"], items[1]["id"], 55)])
    r = await _preview(client, sid, session["id"], h, edited)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["valid"] is True
    assert body["changed_rows"] == 1
    assert body["changed_scores"] == 2


async def test_preview_aggregates_row_errors(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, students=2, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    items = session["items"]
    a0 = session["participants"][0]["attempts"][0]
    a1 = session["participants"][1]["attempts"][0]
    content = await _download(client, sid, session["id"], h)
    edited = _set_cells(
        content,
        [
            (a0["attempt_id"], items[0]["id"], 999),   # 超出满分
            (a0["attempt_id"], items[1]["id"], "abc"),  # 格式非法
            (a1["attempt_id"], items[0]["id"], 12.345),  # 精度超两位
        ],
    )
    r = await _preview(client, sid, session["id"], h, edited)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["valid"] is False
    assert len(body["errors"]) == 3
    messages = " ".join(e["message"] for e in body["errors"])
    assert "0 到" in messages and "格式无效" in messages and "两位小数" in messages


async def test_preview_detects_stale_revision(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    items = session["items"]
    a0 = session["participants"][0]["attempts"][0]
    content = await _download(client, sid, session["id"], h)
    # 导出后有人先改了分:attempt.revision 前移,旧成绩册过期。
    r = await _grade_attempt(client, sid, h, a0, [(items[0]["id"], 10)])
    assert r.status_code == 200
    edited = _set_cells(content, [(a0["attempt_id"], items[0]["id"], 30)])
    r = await _preview(client, sid, session["id"], h, edited)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["valid"] is False
    assert any(e["field"] == "revision" for e in body["errors"])


# --------------------------------------------------------------------------- #
# 应用:历史 + 原子 + 幂等 + 冲突
# --------------------------------------------------------------------------- #
async def test_apply_appends_history_and_refreshes_cache(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, students=2, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    items = session["items"]
    a0 = session["participants"][0]["attempts"][0]
    a1 = session["participants"][1]["attempts"][0]
    content = await _download(client, sid, session["id"], h)
    edited = _set_cells(
        content,
        [
            (a0["attempt_id"], items[0]["id"], 40),
            (a0["attempt_id"], items[1]["id"], 60),
            (a1["attempt_id"], items[0]["id"], 20),
        ],
    )
    r = await _apply(client, sid, session["id"], h, edited, batch_id="batch-1")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["applied_attempts"] == 2
    assert body["changed_rows"] == 2
    assert body["changed_scores"] == 3
    assert body["replayed"] is False

    grade_count = await db_session.scalar(select(func.count()).select_from(ResponseGrade))
    assert grade_count == 3  # 追加式:三条评分历史

    # 缓存刷新:gradebook 里 a0 满分 100。
    r = await client.get(
        f"{API}/subjects/{sid}/assessment-sessions/{session['id']}/gradebook", headers=h
    )
    assert r.status_code == 200
    rows = {row["participation_id"]: row for row in r.json()["participants"]}
    p0 = session["participants"][0]["id"]
    assert Decimal(rows[p0]["total_score"]) == Decimal("100")


async def test_apply_rejects_file_with_errors_atomically(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, students=2, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    items = session["items"]
    a0 = session["participants"][0]["attempts"][0]
    content = await _download(client, sid, session["id"], h)
    # 一处合法 + 一处越界:整文件原子拒绝,不写任何分。
    edited = _set_cells(
        content,
        [(a0["attempt_id"], items[0]["id"], 40), (a0["attempt_id"], items[1]["id"], 999)],
    )
    r = await _apply(client, sid, session["id"], h, edited, batch_id="batch-bad")
    assert r.status_code == 422, r.text
    grade_count = await db_session.scalar(select(func.count()).select_from(ResponseGrade))
    assert grade_count == 0


async def test_apply_idempotent_replay(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    items = session["items"]
    a0 = session["participants"][0]["attempts"][0]
    content = await _download(client, sid, session["id"], h)
    edited = _set_cells(content, [(a0["attempt_id"], items[0]["id"], 40)])

    r1 = await _apply(client, sid, session["id"], h, edited, batch_id="batch-x")
    assert r1.status_code == 200, r1.text
    assert r1.json()["replayed"] is False
    count_after_first = await db_session.scalar(select(func.count()).select_from(ResponseGrade))

    # 同 batch + 同文件重放:返回原结果,不再写库。
    r2 = await _apply(client, sid, session["id"], h, edited, batch_id="batch-x")
    assert r2.status_code == 200, r2.text
    assert r2.json()["replayed"] is True
    assert r2.json()["changed_scores"] == r1.json()["changed_scores"]
    count_after_replay = await db_session.scalar(select(func.count()).select_from(ResponseGrade))
    assert count_after_replay == count_after_first


async def test_apply_same_batch_different_file_conflicts(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    items = session["items"]
    a0 = session["participants"][0]["attempts"][0]
    content = await _download(client, sid, session["id"], h)

    first = _set_cells(content, [(a0["attempt_id"], items[0]["id"], 40)])
    r1 = await _apply(client, sid, session["id"], h, first, batch_id="batch-dup")
    assert r1.status_code == 200, r1.text

    # 相同 batch_id,内容不同(sha256 不符)→ 409。
    second = _set_cells(content, [(a0["attempt_id"], items[0]["id"], 30)])
    r2 = await _apply(client, sid, session["id"], h, second, batch_id="batch-dup")
    assert r2.status_code == 409, r2.text


# --------------------------------------------------------------------------- #
# 权限矩阵
# --------------------------------------------------------------------------- #
async def test_gradebook_permission_matrix(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    content = await _download(client, sid, session["id"], h)

    # viewer(仅 VIEW_QUESTION):导出 / 预览 / 应用均 403。
    hv = _auth(ctx["viewer"])
    assert (await client.get(
        f"{API}/subjects/{sid}/assessment-sessions/{session['id']}/gradebook.xlsx", headers=hv
    )).status_code == 403
    assert (await _preview(client, sid, session["id"], hv, content)).status_code == 403
    assert (await _apply(client, sid, session["id"], hv, content, batch_id="b")).status_code == 403

    # editor(VIEW_ASSESSMENT + EDIT_SCORE):三者均放行。
    he = _auth(ctx["editor"])
    assert (await client.get(
        f"{API}/subjects/{sid}/assessment-sessions/{session['id']}/gradebook.xlsx", headers=he
    )).status_code == 200
    assert (await _preview(client, sid, session["id"], he, content)).status_code == 200
    assert (await _apply(client, sid, session["id"], he, content, batch_id="editor-batch")).status_code == 200


# --------------------------------------------------------------------------- #
# 跨学科 404
# --------------------------------------------------------------------------- #
async def test_cross_subject_404(client, ctx, db_session):
    sid, h, qs, version, session = await _make_session(client, ctx, db_session, scores=(40, 60))
    await _start(client, sid, session["id"], h)
    content = await _download(client, sid, session["id"], h)
    sid2 = ctx["subject2"].id
    hm2 = _auth(ctx["manager"])  # manager 同时是 subject2 的 manager
    assert (await client.get(
        f"{API}/subjects/{sid2}/assessment-sessions/{session['id']}/gradebook.xlsx", headers=hm2
    )).status_code == 404
    assert (await _preview(client, sid2, session["id"], hm2, content)).status_code == 404
    assert (await _apply(client, sid2, session["id"], hm2, content, batch_id="x")).status_code == 404

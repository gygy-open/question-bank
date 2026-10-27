"""
Shared Excel import infrastructure.

Each import declares an ``ImportSchema`` (the single source of truth for its
columns). Header labels, sheet titles, example rows and row-level messages live
in ``EXCEL_TEXT`` keyed by locale. Templates are rendered from schema + locale;
parsing accepts the header of *any* supported locale and maps it back to the
column ``key``, so a Chinese template imports fine under an English UI and vice
versa. Business logic must only use column keys, never header text.
"""
import io
from dataclasses import dataclass
from typing import Dict, List, Optional, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.locale import DEFAULT_LOCALE, SUPPORTED_LOCALES, get_locale, normalize_locale


@dataclass(frozen=True)
class Column:
    key: str
    required: bool = False
    width: int = 16


@dataclass(frozen=True)
class ImportSchema:
    name: str
    columns: Sequence[Column]

    @property
    def keys(self) -> List[str]:
        return [c.key for c in self.columns]

    @property
    def required_keys(self) -> List[str]:
        return [c.key for c in self.columns if c.required]


class MissingColumnsError(ValueError):
    def __init__(self, schema: "ImportSchema", keys: List[str]) -> None:
        self.schema = schema
        self.keys = keys
        super().__init__(", ".join(keys))


USER_SCHEMA = ImportSchema(
    "user",
    [
        Column("username", True, 18),
        Column("full_name", True),
        Column("password", True),
        Column("admin"),
    ],
)
TAG_SCHEMA = ImportSchema(
    "tag",
    [Column("name", True, 20), Column("category"), Column("color")],
)
KP_LEVELS = 8
KP_SCHEMA = ImportSchema(
    "knowledge_point",
    [Column("subject", True)] + [Column(f"level_{i}", True, 18) for i in range(1, KP_LEVELS + 1)],
)

_LEVEL_ZH = ["一", "二", "三", "四", "五", "六", "七", "八"]

EXCEL_TEXT: Dict[str, Dict[str, str]] = {
    "zh-CN": {
        "missing_columns": "缺少必需的列：{columns}。请使用最新导入模板。",
        # user
        "user.sheet": "用户导入",
        "user.col.username": "用户名",
        "user.col.full_name": "姓名",
        "user.col.password": "密码",
        "user.col.admin": "管理员",
        "user.err.username_required": "用户名不能为空",
        "user.err.full_name_required": "姓名不能为空",
        "user.err.password_required": "密码不能为空",
        "user.err.password_too_short": "密码长度不能少于 {min} 位",
        "user.err.duplicate_in_file": "用户名 '{username}' 在文件中重复",
        "user.err.exists": "用户名 '{username}' 已存在",
        # tag
        "tag.sheet": "标签导入",
        "tag.col.name": "名称",
        "tag.col.category": "分类",
        "tag.col.color": "颜色",
        "tag.err.name_required": "名称不能为空",
        "tag.err.bad_color": "颜色 '{color}' 格式不正确，需为十六进制色值（如 #4F81BD）",
        # knowledge point
        "knowledge_point.sheet": "知识点导入",
        "knowledge_point.col.subject": "学科名称",
        **{f"knowledge_point.col.level_{i + 1}": f"{_LEVEL_ZH[i]}级目录" for i in range(KP_LEVELS)},
        "knowledge_point.err.subject_required": "学科名称不能为空",
        "knowledge_point.err.subject_not_found": "学科 '{subject}' 不存在，请先创建该学科",
        "knowledge_point.err.level_1_required": "一级目录不能为空",
    },
    "en-US": {
        "missing_columns": "Missing required columns: {columns}. Please use the latest import template.",
        "user.sheet": "User Import",
        "user.col.username": "Username",
        "user.col.full_name": "Full Name",
        "user.col.password": "Password",
        "user.col.admin": "Admin",
        "user.err.username_required": "Username is required",
        "user.err.full_name_required": "Full name is required",
        "user.err.password_required": "Password is required",
        "user.err.password_too_short": "Password must be at least {min} characters",
        "user.err.duplicate_in_file": "Username '{username}' is duplicated in the file",
        "user.err.exists": "Username '{username}' already exists",
        "tag.sheet": "Tag Import",
        "tag.col.name": "Name",
        "tag.col.category": "Category",
        "tag.col.color": "Color",
        "tag.err.name_required": "Name is required",
        "tag.err.bad_color": "Color '{color}' is invalid; use a hex color such as #4F81BD",
        "knowledge_point.sheet": "Knowledge Point Import",
        "knowledge_point.col.subject": "Subject Name",
        **{f"knowledge_point.col.level_{i + 1}": f"Level {i + 1}" for i in range(KP_LEVELS)},
        "knowledge_point.err.subject_required": "Subject name is required",
        "knowledge_point.err.subject_not_found": "Subject '{subject}' does not exist; create it first",
        "knowledge_point.err.level_1_required": "Level 1 is required",
    },
}

# Example rows shipped in templates. Keyed by schema name, then locale.
EXCEL_EXAMPLES: Dict[str, Dict[str, List[List[str]]]] = {
    "user": {
        "zh-CN": [
            ["teacher01", "张三", "changeme123", ""],
            ["admin02", "李四", "changeme123", "是"],
        ],
        "en-US": [
            ["teacher01", "Alex Smith", "changeme123", ""],
            ["admin02", "Jamie Lee", "changeme123", "Yes"],
        ],
    },
    "tag": {
        "zh-CN": [["高考真题", "来源", "#4F81BD"], ["易错题", "", ""]],
        "en-US": [["Past Exam", "Source", "#4F81BD"], ["Common Mistakes", "", ""]],
    },
    "knowledge_point": {
        "zh-CN": [
            ["数学", "代数", "方程", "一元一次方程"],
            ["数学", "代数", "方程", "一元二次方程", "解法", "公式法"],
            ["数学", "几何", "平面几何", "三角形"],
        ],
        "en-US": [
            ["Math", "Algebra", "Equations", "Linear Equations"],
            ["Math", "Algebra", "Equations", "Quadratic Equations", "Methods", "Quadratic Formula"],
            ["Math", "Geometry", "Plane Geometry", "Triangles"],
        ],
    },
}

# Values accepted as "true" for boolean cells, regardless of UI language.
TRUE_VALUES = {"是", "true", "1", "yes", "y", "admin", "管理员"}


def text(key: str, locale: Optional[str] = None, **params: object) -> str:
    loc = normalize_locale(locale) if locale else get_locale()
    template = EXCEL_TEXT.get(loc, EXCEL_TEXT[DEFAULT_LOCALE]).get(key)
    if template is None:
        template = EXCEL_TEXT[DEFAULT_LOCALE][key]
    return template.format(**params) if params else template


def header(schema: ImportSchema, key: str, locale: Optional[str] = None) -> str:
    return text(f"{schema.name}.col.{key}", locale)


def _aliases(schema: ImportSchema) -> Dict[str, str]:
    """Map every locale's header label (casefolded) to its column key."""
    mapping: Dict[str, str] = {}
    for loc in SUPPORTED_LOCALES:
        for col in schema.columns:
            mapping[header(schema, col.key, loc).strip().casefold()] = col.key
    return mapping


def generate_template(schema: ImportSchema, locale: Optional[str] = None) -> bytes:
    loc = normalize_locale(locale) if locale else get_locale()
    wb = Workbook()
    ws = wb.active
    ws.title = text(f"{schema.name}.sheet", loc)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

    ws.append([header(schema, c.key, loc) for c in schema.columns])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    width = len(schema.columns)
    for row in EXCEL_EXAMPLES[schema.name][loc]:
        ws.append(list(row) + [""] * (width - len(row)))

    for idx, col in enumerate(schema.columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = col.width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def read_rows(
    file_bytes: bytes, schema: ImportSchema, required: Optional[Sequence[str]] = None
) -> pd.DataFrame:
    """Read the .xlsx and rename headers (any locale) to column keys."""
    df = pd.read_excel(io.BytesIO(file_bytes), dtype=str)
    aliases = _aliases(schema)
    df = df.rename(columns=lambda c: aliases.get(str(c).strip().casefold(), c))
    needed = list(required) if required is not None else schema.required_keys
    missing = [k for k in needed if k not in df.columns]
    if missing:
        raise MissingColumnsError(schema, missing)
    return df


def missing_columns_message(exc: MissingColumnsError) -> str:
    labels = [header(exc.schema, k) for k in exc.keys]
    return text("missing_columns", columns=", ".join(labels))


def clean(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()

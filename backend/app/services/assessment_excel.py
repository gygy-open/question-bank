"""通用评测成绩册 (Assessment gradebook) 的 Excel 传输 —— v2 契约。

可见工作表面向教师:参与者 identifier / display_name、每个 AssessmentItem 一列、总分。
隐藏 `_meta` 工作表用稳定且版本化的二维结构把工作簿绑定到一次投放(session),并保留
条目 / 作答 ID 与乐观锁 revision —— 解析时**只依赖 `_meta` 的结构化 ID,绝不反解显示文本**。

v2 与旧的 Exam 版本不兼容:模型改为 Assessment* 全新契约(session / item / attempt /
ResponseGrade 追加式评分),`_meta` 显式携带 format / version / session_id / 条目列 /
行映射,以便未来演进时能安全拒绝旧文件。
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.assessment import (
    SCORE_SCALE,
    AssessmentResponse,
    AssessmentSession,
    GradeStatus,
    ResponseGrade,
)

# 可见工作表与隐藏元数据工作表名。
SHEET_NAME = "成绩录入"
META_SHEET_NAME = "_meta"

# 版本化契约标识:改结构必须升 version,旧文件据此被安全拒绝。
FORMAT = "question-bank-assessment-gradebook"
FORMAT_VERSION = 2

# 可见表首两列固定为标识 / 姓名,条目列从第 3 列开始。
_FIRST_ITEM_COL = 3


@dataclass(frozen=True)
class ScoreImportIssue:
    """单行级别的解析错误(可聚合,不中断整体解析)。"""
    row: int
    field: str
    message: str


@dataclass(frozen=True)
class ImportRow:
    """一条可应用的导入行:绑定 attempt(乐观锁 revision)与逐条目分值。"""
    excel_row: int
    participation_id: int
    attempt_id: int
    attempt_revision: int
    scores: Dict[int, Optional[Decimal]]


@dataclass(frozen=True)
class ParsedImport:
    rows: List[ImportRow]
    changed_rows: int
    changed_scores: int
    unchanged_rows: int
    errors: List[ScoreImportIssue] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# 工具
# --------------------------------------------------------------------------- #
def _fmt(value: Decimal) -> str:
    return format(value, "f")


def _field_label(position: int, item_key: str) -> str:
    return f"第 {position + 1} 题"


def _latest_grade(response: AssessmentResponse) -> Optional[ResponseGrade]:
    if not response.grades:
        return None
    return max(response.grades, key=lambda g: g.revision)


def _effective_score(response: Optional[AssessmentResponse]) -> Optional[Decimal]:
    """当前最新 ResponseGrade 的有效得分;未评分 / 无分返回 None(单元格留空)。"""
    if response is None:
        return None
    latest = _latest_grade(response)
    if latest is None or latest.status != GradeStatus.GRADED.value or latest.score is None:
        return None
    return latest.score


def _decimal(value: Any) -> Optional[Decimal]:
    """单元格值 → Decimal;空白返回 None;布尔 / 非法格式抛异常交由调用方转成行错误。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise InvalidOperation
    return Decimal(str(value).strip())


def _as_int(value: Any) -> Optional[int]:
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------------- #
# 导出:教师成绩册
# --------------------------------------------------------------------------- #
def generate_gradebook(session: AssessmentSession) -> bytes:
    """生成教师成绩册工作簿:可见表 + 隐藏 `_meta`(结构化绑定)。"""
    items = sorted(session.version.items, key=lambda it: it.position)
    participations = list(session.participations)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    meta = workbook.create_sheet(META_SHEET_NAME)

    total_col = _FIRST_ITEM_COL + len(items)  # 1-based:总分列
    headers: List[str] = ["标识", "姓名"]
    for position, item in enumerate(items):
        headers.append(f"{position + 1}. {item.item_key}（满分 {_fmt(item.max_score)}）")
    headers.append("总分")
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="2563EB")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "C2"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 16
    for column in range(_FIRST_ITEM_COL, total_col + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 16

    # --- 隐藏 _meta:稳定、版本化的二维结构 ---
    meta.append(["format", FORMAT])
    meta.append(["version", FORMAT_VERSION])
    meta.append(["session_id", session.id])
    meta.append(["items"])
    meta.append(["column", "item_id", "item_key", "max_score"])
    for offset, item in enumerate(items):
        meta.append([_FIRST_ITEM_COL + offset, item.id, item.item_key, _fmt(item.max_score)])
    meta.append(["rows"])
    meta.append(["excel_row", "participation_id", "attempt_id", "attempt_revision"])

    first_letter = get_column_letter(_FIRST_ITEM_COL)
    last_letter = get_column_letter(total_col - 1)
    excel_row = 2
    for participation in participations:
        attempt = participation.attempts[-1] if participation.attempts else None
        responses = {r.item_id: r for r in attempt.responses} if attempt is not None else {}
        values: List[Any] = [participation.identifier, participation.display_name]
        for item in items:
            values.append(_effective_score(responses.get(item.id)))
        if items:
            values.append(
                f"=IF(COUNT({first_letter}{excel_row}:{last_letter}{excel_row})={len(items)},"
                f'SUM({first_letter}{excel_row}:{last_letter}{excel_row}),"")'
            )
        else:
            values.append(None)
        sheet.append(values)
        for column in range(_FIRST_ITEM_COL, total_col + 1):
            sheet.cell(excel_row, column).number_format = "0.00"
        meta.append(
            [
                excel_row,
                participation.id,
                attempt.id if attempt is not None else None,
                attempt.revision if attempt is not None else None,
            ]
        )
        excel_row += 1

    meta.sheet_state = "hidden"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


# --------------------------------------------------------------------------- #
# 解析:成绩册 → 结构化导入
# --------------------------------------------------------------------------- #
def _read_meta_head(meta_rows: List[tuple]) -> Dict[str, Any]:
    head: Dict[str, Any] = {}
    for row in meta_rows:
        if not row:
            continue
        key = row[0]
        if key in ("format", "version", "session_id"):
            head[key] = row[1] if len(row) > 1 else None
    return head


def _read_meta_sections(meta_rows: List[tuple]) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """从 `_meta` 解析条目列与行映射两段结构。"""
    item_cols: List[Dict[str, Any]] = []
    row_maps: List[Dict[str, Any]] = []
    section: Optional[str] = None
    for row in meta_rows:
        if not row:
            continue
        marker = row[0]
        if marker == "items":
            section = "items_header"
            continue
        if marker == "rows":
            section = "rows_header"
            continue
        if section == "items_header":
            section = "items"
            continue
        if section == "rows_header":
            section = "rows"
            continue
        if section == "items":
            item_id = _as_int(row[1] if len(row) > 1 else None)
            if item_id is None:
                continue
            item_cols.append(
                {
                    "column": _as_int(row[0]),
                    "item_id": item_id,
                    "item_key": row[2] if len(row) > 2 else None,
                    "max_score": row[3] if len(row) > 3 else None,
                }
            )
        elif section == "rows":
            excel_row = _as_int(row[0])
            if excel_row is None:
                continue
            row_maps.append(
                {
                    "excel_row": excel_row,
                    "participation_id": _as_int(row[1] if len(row) > 1 else None),
                    "attempt_id": _as_int(row[2] if len(row) > 2 else None),
                    "attempt_revision": _as_int(row[3] if len(row) > 3 else None),
                }
            )
    return item_cols, row_maps


def parse_gradebook(file_bytes: bytes, session: AssessmentSession) -> ParsedImport:
    """解析上传的成绩册。

    结构级失败(不可读 / 缺表 / format/version/session 不符 / 条目结构与评测不一致)抛
    DomainError;行级问题(未知/重复 attempt、stale revision、分值格式/范围/精度)聚合为
    错误列表返回,绝不写库。
    """
    # 延迟导入以打破 capabilities → services → excel 的模块级循环。
    from app.capabilities.errors import Invalid, Unprocessable

    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=False)
    except Exception as exc:  # openpyxl 抛的异常类型不稳定,统一转成领域错误。
        raise Invalid("无法读取 Excel 文件,请使用系统导出的成绩册模板") from exc

    if SHEET_NAME not in workbook.sheetnames or META_SHEET_NAME not in workbook.sheetnames:
        raise Invalid("成绩册缺少必要工作表,请重新导出")

    sheet = workbook[SHEET_NAME]
    meta = workbook[META_SHEET_NAME]
    meta_rows = list(meta.iter_rows(values_only=True))

    head = _read_meta_head(meta_rows)
    if head.get("format") != FORMAT:
        raise Invalid("成绩册格式不受支持,请重新导出")
    if _as_int(head.get("version")) != FORMAT_VERSION:
        raise Invalid("成绩册版本不受支持,请重新导出")
    if _as_int(head.get("session_id")) != session.id:
        raise Invalid("成绩册不属于当前评测")

    item_cols, row_maps = _read_meta_sections(meta_rows)

    items = sorted(session.version.items, key=lambda it: it.position)
    item_by_id = {it.id: it for it in items}
    meta_item_ids = [c["item_id"] for c in item_cols]
    if len(meta_item_ids) != len(set(meta_item_ids)):
        raise Unprocessable("成绩册题目结构损坏:条目列重复")
    if set(meta_item_ids) != set(item_by_id):
        raise Unprocessable("成绩册题目结构与当前评测不一致,请重新导出")
    for col in item_cols:
        if col["column"] is None:
            raise Unprocessable("成绩册题目结构损坏:缺少条目列号")
    if not row_maps:
        raise Unprocessable("成绩册缺少参与者行映射,请重新导出")

    attempts_by_id: Dict[int, Any] = {}
    for participation in session.participations:
        for attempt in participation.attempts:
            attempts_by_id[attempt.id] = attempt

    errors: List[ScoreImportIssue] = []
    rows: List[ImportRow] = []
    changed_rows = 0
    changed_scores = 0
    seen_attempts: set[int] = set()

    for rm in row_maps:
        excel_row = rm["excel_row"]
        attempt_id = rm["attempt_id"]
        participation_id = rm["participation_id"]
        expected_revision = rm["attempt_revision"]

        if attempt_id is None:
            # 无作答的参与者行(理论上不出现)——无从录分,直接跳过。
            continue
        if attempt_id in seen_attempts:
            errors.append(ScoreImportIssue(excel_row, "作答", "成绩册存在重复的作答行"))
            continue
        seen_attempts.add(attempt_id)

        attempt = attempts_by_id.get(attempt_id)
        if attempt is None or (
            participation_id is not None and attempt.participation_id != participation_id
        ):
            errors.append(ScoreImportIssue(excel_row, "作答", "作答不属于当前评测"))
            continue
        if expected_revision is None or attempt.revision != expected_revision:
            errors.append(ScoreImportIssue(excel_row, "revision", "成绩已被修改,请重新导出成绩册"))
            continue

        existing: Dict[int, Optional[Decimal]] = {}
        for response in attempt.responses:
            existing[response.item_id] = _effective_score(response)

        parsed_scores: Dict[int, Optional[Decimal]] = {}
        row_changes = 0
        row_has_error = False
        for col in item_cols:
            item = item_by_id[col["item_id"]]
            label = _field_label(item.position, item.item_key)
            try:
                score = _decimal(sheet.cell(excel_row, col["column"]).value)
            except (InvalidOperation, ValueError):
                errors.append(ScoreImportIssue(excel_row, label, "分值格式无效"))
                row_has_error = True
                continue
            if score is not None:
                if score < 0 or score > item.max_score:
                    errors.append(
                        ScoreImportIssue(
                            excel_row, label, f"分值必须在 0 到 {_fmt(item.max_score)} 之间"
                        )
                    )
                    row_has_error = True
                    continue
                exponent = score.as_tuple().exponent
                if isinstance(exponent, int) and exponent < -SCORE_SCALE:
                    errors.append(ScoreImportIssue(excel_row, label, "分值精度最多两位小数"))
                    row_has_error = True
                    continue
            parsed_scores[item.id] = score
            if existing.get(item.id) != score:
                row_changes += 1

        if row_has_error:
            continue
        rows.append(
            ImportRow(
                excel_row=excel_row,
                participation_id=attempt.participation_id,
                attempt_id=attempt_id,
                attempt_revision=expected_revision,
                scores=parsed_scores,
            )
        )
        if row_changes:
            changed_rows += 1
            changed_scores += row_changes

    unchanged_rows = len(rows) - changed_rows
    return ParsedImport(
        rows=rows,
        changed_rows=changed_rows,
        changed_scores=changed_scores,
        unchanged_rows=unchanged_rows,
        errors=errors,
    )
